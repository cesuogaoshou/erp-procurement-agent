"""Build the final course report, screenshots evidence, and source package."""
from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(__file__).resolve().parents[1]
DELIVERY = ROOT / "delivery" / "A+B ERP课程作业报告"
DOCX_NAME = "A+B ERP课程作业报告.docx"
SCREENSHOTS = ROOT / "reports/generated/screenshots_current"
CHARTS = ROOT / "reports/generated/report_charts"
EVIDENCE = ROOT / "reports/generated"
DATA = ROOT / "data/raw"


def load(name: str):
    return json.loads((EVIDENCE / name).read_text(encoding="utf-8"))


def set_east_asia(font, name: str) -> None:
    font.name = name
    font._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    set_east_asia(normal.font, "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "黑体"
        style.font.size = Pt(size)
        style.font.bold = True
        set_east_asia(style.font, "黑体")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def add_page_number_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fld_run = OxmlElement("w:r")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_run.append(fld_text)
    fld.append(fld_run)
    paragraph._p.append(fld)
    paragraph.add_run(" 页").font.size = Pt(9)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, *, bold: bool = False, size: float | None = None, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.line_spacing = 1.4


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.line_spacing = 1.4


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], *, widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for index, header in enumerate(headers):
        hdr[index].text = ""
        p = hdr[index].paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(9.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(hdr[index], "E7EEF8")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            p = cells[index].paragraphs[0]
            run = p.add_run(str(value))
            run.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(0)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)
    p.paragraph_format.space_after = Pt(8)


def add_picture(doc: Document, path: Path, caption: str, width: float = 15.0) -> None:
    if not path.exists():
        add_para(doc, f"[缺少图片：{path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width))
    add_caption(doc, caption)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = "Consolas"
    set_east_asia(run.font, "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(40, 40, 40)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_report() -> Path:
    baseline = load("baseline_eval.json")
    forecast = load("forecast_comparison.json")["models"]
    supplier = load("supplier_eval.json")
    agent = load("agent_input_eval.json")["summary"]
    mysql = load("mysql_integration.json")

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    add_page_number_footer(doc)

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ERP 智能采购 Agent")
    run.font.name = "黑体"
    set_east_asia(run.font, "黑体")
    run.font.size = Pt(28)
    run.bold = True
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("课程大作业报告（算法研究类）")
    run2.font.name = "黑体"
    set_east_asia(run2.font, "黑体")
    run2.font.size = Pt(18)
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("基于大语言模型与混合需求预测的\nERP 智能采购决策系统")
    run.font.name = "宋体"
    set_east_asia(run.font, "宋体")
    run.font.size = Pt(16)
    run.bold = True
    for _ in range(5):
        doc.add_paragraph()
    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("组员 A", "________________"),
        ("组员 B", "________________"),
        ("学号 / 班级", "________________"),
        ("完成日期", "2026 年 10 月 30 日"),
    ]
    for row, (label, value) in zip(info.rows, values):
        row.cells[0].text = label
        row.cells[1].text = value
        for cell in row.cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(11)
    page_break(doc)

    # Abstract
    add_heading(doc, "摘要", 1)
    add_para(doc, "本报告面向课程选题 3.2“采购订单自动生成与供应商推荐算法研究”，设计并实现了一个 ERP 智能采购决策系统。系统以 Python/FastAPI 为后端、Vue 3 为前端，围绕物料库存、需求预测、缺口计算、供应商推荐、采购草稿和人工审批构建完整业务闭环。LLM 仅负责自然语言理解与 Function Calling 工具编排，SQL 查询、数值计算、评分、预算约束、审批权限和事务幂等全部由后端服务完成。")
    add_para(doc, "需求预测采用“动态 Prophet 主路径 + 30 天移动平均兜底”的混合策略；供应商推荐综合历史价格、准时交付率、质量合格率和投诉率四项指标，先做启用状态、供货关系、MOQ、交期、产能和预算过滤，再做 Min-Max 归一化和加权评分，输出 Top-3。系统同时实现了采购草稿证据链、整单预算硬约束、草稿人工编辑、审批版本校验、MySQL 行锁并发控制和 Agent 输入评测。")
    add_para(doc, "实验结果表明：在 160 个 30 天滚动窗口上，混合策略的缺口 MAPE 为 5.72%，动态 Prophet 在 140 个成功窗口上的缺口 MAPE 为 5.34%；供应商综合评分在冻结集和挑战集上的 Top-3 命中率均为 100%，高于价格排序基线；26 条真实模型 Agent 输入用例全部通过。报告最后分析了固定库存快照、测试窗口复看、供应商标签模拟等局限，并给出“候选集 + 模型选 ID + 人工确认”的后续改进方向。")
    p = doc.add_paragraph()
    p.add_run("关键词：").bold = True
    p.add_run("ERP 采购；需求预测；Prophet；供应商推荐；Function Calling；LangGraph；FastAPI")
    page_break(doc)

    # TOC
    add_heading(doc, "目录", 1)
    toc_items = [
        "1 项目背景与课程要求",
        "2 需求分析与系统范围",
        "3 数据集定义与预处理",
        "4 系统架构与实现过程",
        "5 算法介绍与改进",
        "6 实验过程与截图",
        "7 实验结果分析",
        "8 结论与未来工作",
        "附录 A 复现命令与项目结构",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.line_spacing = 1.5
    page_break(doc)

    # 1
    add_heading(doc, "1 项目背景与课程要求", 1)
    add_heading(doc, "1.1 项目背景", 2)
    add_para(doc, "制造企业的采购计划通常依赖人工汇总库存、历史消耗、供应商报价和交期。数据分散、口径不统一时，容易出现缺料、过量采购和审批依据不足。大语言模型可以降低自然语言操作门槛，但采购金额、供应商资格和审批权限不能交给模型自由决定。本项目因此采用“LLM 理解与编排、后端计算与校验”的分层架构。")
    add_heading(doc, "1.2 课程要求对应", 2)
    add_table(doc, ["课程要求", "本系统实现与证据"], [
        ["LLM 理解自然语言采购指令，Function Calling 编排工作流", "OpenAI-compatible Function Calling + 有界 Agent Loop + 持久化 tool trace"],
        ["先查库存，再比较价格、交期、质量，推荐供应商后生成订单", "inventory → forecast → recommend_suppliers → create_purchase_draft 工具链"],
        ["供应商评分综合历史价格、准时交付、质量与投诉", "四指标过滤 + Min-Max 归一化 + 0.35/0.30/0.25/0.10 加权评分"],
        ["物料缺口预测采用 Prophet/LSTM", "动态 Prophet 主路径；移动平均兜底；不使用 LSTM"],
        ["公开 ERP 数据或自建模拟 ERP 数据库", "固定种子 20261010 的模拟数据 + CSV + 数据字典"],
        ["缺口预测 MAPE 低于 8%", "160 个 30 天滚动窗口，混合策略缺口 MAPE 5.72%"],
        ["供应商 Top-3 命中率不低于 85%", "冻结集/挑战集综合评分 Top-3 均为 100%"],
        ["权限、审批、日志、业务规则与人工确认", "buyer/approver 角色、版本校验、幂等、审计日志、审批事务"],
        ["自主贡献", "自建数据、混合预测策略、证据链草稿、供应商评测与 Agent 输入评测"],
        ["Python/Java 源码、数据集、至少 4 页 Word 报告", "本报告 + 源代码和数据集目录"],
    ], widths=[7.3, 8.0])
    add_heading(doc, "1.3 主要工作", 2)
    add_bullets(doc, [
        "完成 10 物料、10 供应商、540 天需求的模拟 ERP 数据与可复现导出。",
        "实现库存、预测、缺口、供应商推荐、采购草稿和审批的确定性后端。",
        "实现 Function Calling Agent、LangGraph 执行、工具调用 Trace 持久化。",
        "实现 Vue 3 四页面工作台与草稿数量/预算编辑。",
        "完成预测、推荐、Agent、MySQL 并发和审批安全实验。",
    ])
    page_break(doc)

    # 2
    add_heading(doc, "2 需求分析与系统范围", 1)
    add_heading(doc, "2.1 功能需求", 2)
    add_bullets(doc, [
        "物料与库存查询：按物料编码/名称/别名查询库存、预留和可用量。",
        "需求预测与缺口：按预测窗口输出累计需求、安全库存、可用库存和缺口。",
        "供应商推荐：按约束过滤并用四项指标评分，返回 Top-3 及被排除原因。",
        "采购草稿：服务端校验预测/推荐证据、数量、MOQ、包装倍数、产能和整单预算。",
        "人工审批：采购员提交，审批员批准或驳回，生成正式订单号，记录审批与审计日志。",
        "Agent 编排：自然语言到工具调用，审批工具不暴露给模型。",
        "前端工作台：智能采购、数据分析、采购审批、执行记录四个页面。",
    ])
    add_heading(doc, "2.2 非功能需求", 2)
    add_bullets(doc, [
        "正确性：金额、数量、评分、缺口由后端 Decimal 计算，模型不能直接写库。",
        "安全性：服务端签发 Bearer Token；角色从认证上下文取得；审批动作不接受模型调用。",
        "一致性：审批使用数据库事务、行锁、版本号和幂等键。",
        "可复现性：固定随机种子、CSV 导出、迁移脚本、评测脚本和原始 JSON。",
        "可解释性：返回推荐依据、预测模型、降级原因和证据引用。",
    ])
    add_heading(doc, "2.3 范围边界", 2)
    add_para(doc, "系统不实现真实外部供应商系统、支付、邮件、完整 ERP 模块、RAG、MCP 服务、多 Agent 或 LSTM。模拟供应商和模拟数据仅用于课程演示，不能表述为真实生产数据。")
    page_break(doc)

    # 3
    add_heading(doc, "3 数据集定义与预处理", 1)
    add_heading(doc, "3.1 数据规模", 2)
    add_table(doc, ["数据表", "规模", "说明"], [
        ["materials", "10 行", "物料编码、名称、单位、安全库存"],
        ["inventory", "10 行", "现有库存、预留、版本、快照日期"],
        ["demand_history", "5,400 行", "10 物料 × 540 天日需求"],
        ["suppliers", "10 行", "供应商名称与启用状态"],
        ["supplier_materials", "100 行", "报价、交期、MOQ、包装倍数、产能"],
        ["supplier_performance", "100 行", "准时交付、质量检验、投诉记录"],
        ["material_aliases", "35 行", "简称/别名，用于数据库驱动的物料解析"],
        ["data_splits", "3 行", "训练、验证、测试日期范围"],
    ], widths=[4.2, 2.4, 8.8])
    add_heading(doc, "3.2 数据生成与预处理", 2)
    add_para(doc, "数据生成使用固定随机种子 20261010，基准日为 2026-10-10，历史长度为 540 天。需求模式包含趋势、周季节性、年度季节波动、噪声和偶发尖峰。CSV 以 UTF-8 导出，生成元数据记录种子、物料数、供应商数和库存快照日期。")
    add_table(doc, ["划分", "起止日期", "用途"], [
        ["train", "2025-04-18 至 2026-06-11", "模型训练/滚动窗口训练段"],
        ["validation", "2026-06-12 至 2026-08-10", "参数与策略验证"],
        ["test", "2026-08-11 至 2026-10-09", "探索性留出评估"],
    ], widths=[3.0, 7.5, 5.0])
    add_para(doc, "预处理步骤：按物料和日期排序；缺失值不填充伪数据；金额和数量统一使用 Decimal；预测只使用截止日前的历史，未来窗口用于评估；库存使用 2026-10-10 固定快照；Phase 1 不建模已确认在途库存（inbound_confirmed 固定为 0）。")
    add_para(doc, "需要注意：测试窗口在诊断阶段已被查看，因此本报告将其表述为“探索性滚动回测”，不声称是完全未触碰的最终独立测试集；缺口评估使用固定库存快照，不是历史库存回放。")
    page_break(doc)

    # 4
    add_heading(doc, "4 系统架构与实现过程", 1)
    add_heading(doc, "4.1 总体架构", 2)
    add_para(doc, "系统分为 API 层、业务服务层、算法层、数据模型层、Agent 层和前端展示层。请求从 Vue 页面进入 FastAPI，路由完成认证和 Pydantic 校验后调用服务层；服务层通过 SQLAlchemy 访问 SQLite/MySQL；Agent 层只负责选择工具和填写参数，工具执行仍回到服务层。")
    add_table(doc, ["层次", "主要技术", "职责"], [
        ["前端", "Vue 3 + TypeScript + Element Plus + ECharts", "四页面工作台、图表、审批与明细编辑"],
        ["API", "FastAPI + Pydantic", "路由、鉴权、参数校验、错误映射"],
        ["业务服务", "Python 服务类", "库存、预测、推荐、草稿、审批、审计、幂等"],
        ["算法", "Prophet、NumPy/Pandas、Decimal", "需求预测、缺口、归一化和评分"],
        ["数据", "SQLAlchemy 2.0 + Alembic + SQLite/MySQL", "ORM、迁移、事务、行锁、测试数据库"],
        ["Agent", "OpenAI-compatible Function Calling + LangGraph", "工具选择、参数生成、Trace 持久化"],
    ], widths=[2.8, 6.2, 7.2])
    add_heading(doc, "4.2 数据模型与接口", 2)
    add_para(doc, "数据库包含 13 张表：materials、material_aliases、inventory、demand_history、suppliers、supplier_materials、supplier_performance、purchase_orders、purchase_order_items、approval_records、audit_logs、agent_runs、tool_traces。Alembic 共 8 个迁移版本，支持 SQLite 开发和 MySQL 集成。")
    add_para(doc, "后端提供 17 个 REST 接口，包括登录、物料、库存、预测、推荐、草稿、草稿编辑、提交、批准、驳回、Agent run、trace、订单和需求历史。所有金额与数量在 Pydantic/Decimal 层校验，关键写操作由服务层再次校验。")
    add_heading(doc, "4.3 Agent 工具链", 2)
    add_para(doc, "Agent 使用有界循环（最多 8 次工具调用、180 秒超时）。工具包括 capture_request、list_materials、search_materials、resolve_material、get_inventory、get_demand_forecast、recommend_suppliers、create_purchase_draft。审批和驳回工具不在模型工具集中。")
    add_code(doc, "用户文字 → capture_request（物料/数量/预算/交期）\n→ get_demand_forecast（需求、缺口、库存证据）\n→ recommend_suppliers（过滤 + 评分 + Top-3）\n→ create_purchase_draft（服务端证据/预算/MOQ/产能复核）")
    add_heading(doc, "4.4 实现步骤", 2)
    add_numbered(doc, [
        "第一阶段：建立 SQLAlchemy 模型、Alembic 迁移、固定种子模拟数据和 CSV 导出。",
        "第二阶段：实现库存、预测、缺口、供应商推荐、草稿和审批服务及 API。",
        "第三阶段：接入 OpenAI-compatible Function Calling、Agent Loop、LangGraph 和 Trace。",
        "第四阶段：实现 MySQL 事务/行锁/幂等测试、预测与推荐评测、Agent 输入评测。",
        "第五阶段：实现 Vue 四页面、双图表、中文提示、草稿编辑和课程交付包。",
    ])
    page_break(doc)

    # 5
    add_heading(doc, "5 算法介绍与改进", 1)
    add_heading(doc, "5.1 需求预测算法", 2)
    add_para(doc, "需求预测采用 Prophet 加法模型。Prophet 将时间序列分解为趋势、季节性和误差项：")
    add_code(doc, "y(t) = g(t) + s(t) + h(t) + ε(t)\n\ng(t)：分段线性趋势与 changepoint\ns(t)：周/年 Fourier 季节性\nh(t)：节假日项（本项目未配置）\nε(t)：误差项")
    add_para(doc, "本项目关闭日季节性，开启周季节性；历史不少于 365 天时开启年季节性，否则关闭，以避免短历史下年周期不可识别。趋势灵活度 changepoint_prior_scale=0.05，季节性强度 seasonality_prior_scale=10。未来 N 天 yhat 截断负值后求和，得到累计预测需求。")
    add_table(doc, ["参数", "取值", "含义"], [
        ["min_train_days", "90", "少于 90 天不使用 Prophet"],
        ["daily_seasonality", "False", "不拟合日周期"],
        ["weekly_seasonality", "True", "拟合周周期"],
        ["yearly_seasonality", "True（≥365 天）", "历史不足一年时关闭"],
        ["changepoint_prior_scale", "0.05", "趋势转折灵活度"],
        ["seasonality_prior_scale", "10.0", "季节性强度"],
        ["train_window_days", "None", "使用全部历史"],
    ], widths=[5.0, 4.2, 7.0])
    add_para(doc, "改进点：使用“动态 Prophet + 30 天移动平均兜底”的混合策略。历史不足、Prophet/CmdStan 不可用或运行异常时，自动使用最近 30 天平均日需求 × 预测天数，并在返回结果中记录 selected_model_name 和 fallback_reason。")
    add_heading(doc, "5.2 缺口计算", 2)
    add_para(doc, "可用库存和缺口公式如下：")
    add_code(doc, "可用库存 = on_hand - reserved + inbound_confirmed\n缺口 = max(0, 预测需求 + 安全库存 - 可用库存)")
    add_para(doc, "当前 inbound_confirmed 固定为 0；安全库存来自物料主数据，不由模型计算。缺口是计划缺口，不等于最终采购数量；最终数量还要经过供应商产能、MOQ、包装倍数、报价和整单预算校验。")
    add_heading(doc, "5.3 供应商推荐算法", 2)
    add_para(doc, "推荐先过滤：供应商启用、供应关系存在、数量≥MOQ、交期≤要求交期、产能≥数量、金额≤预算。对通过过滤的候选做 Min-Max 归一化：")
    add_code(doc, "成本型（价格、投诉率）：(max - value) / (max - min)\n收益型（交付率、质量率）：(value - min) / (max - min)\n极差为 0 时记 1")
    add_para(doc, "综合评分为 0.35×价格 + 0.30×交付 + 0.25×质量 + 0.10×投诉，按综合分降序、单价升序、供应商 ID 升序取 Top-3。改进点是将“价格排序”作为基线，增加交期、产能、MOQ、预算约束和四项指标评分，避免只看最低价。")
    add_heading(doc, "5.4 Agent 语义解析与安全改进", 2)
    add_para(doc, "Agent 不把物料清单写进提示词。capture_request 从用户原文提取物料查询、数量、预算和交期；后端通过 material_aliases 和候选搜索解析物料 ID。search_materials 返回候选、分数和匹配类型，resolve_material 返回 resolved/ambiguous/not_found。别名必须作为独立词出现，不能把“火星牌超导轴承”中的“轴承”替换成“标准轴承”。")
    add_para(doc, "草稿创建前，后端重新校验预测证据、推荐证据、数量、MOQ/包装、产能、报价和整单预算；草稿可人工修改数量/预算，保存时重新生成证据并记录审计。审批阶段再次校验版本、报价、产能、金额和预算，保证模型输出不能绕过业务规则。")
    page_break(doc)

    # 6
    add_heading(doc, "6 实验过程与截图", 1)
    add_heading(doc, "6.1 实验环境与命令", 2)
    add_table(doc, ["项目", "配置"], [
        ["操作系统", "Windows 10/11"],
        ["Python", "3.13，FastAPI、SQLAlchemy、Prophet、LangGraph"],
        ["前端", "Node.js、Vue 3、Vite、Element Plus、ECharts"],
        ["数据库", "SQLite 演示库；MySQL 8.4 用于集成测试"],
        ["模型", "OpenAI-compatible Function Calling，DeepSeek 模型"],
        ["验证", "pytest、评测脚本、真实模型验收、MySQL 集成测试"],
    ], widths=[4.0, 11.5])
    add_code(doc, "python -m pytest backend/tests -q\npython -m backend.evaluation.run_baseline\npython -m backend.evaluation.run_supplier_eval\npython -m backend.evaluation.run_forecast_comparison\npython -m backend.evaluation.run_forecast_diagnostics\npython -m backend.evaluation.run_agent_input_eval\nnpm --prefix frontend run build")

    add_heading(doc, "6.2 智能采购页面", 2)
    add_para(doc, "采购员在页面输入自然语言需求和整单预算，点击“分析并准备采购”。Agent 依次调用 capture_request、get_demand_forecast、recommend_suppliers、create_purchase_draft，页面显示业务结论、下一步和草稿明细。")
    add_picture(doc, SCREENSHOTS / "01-agent-procurement.png", "图 6-1 智能采购页面：自然语言输入与整单预算", 15.2)
    add_picture(doc, SCREENSHOTS / "02-agent-result.png", "图 6-2 Agent 真实调用结果：草稿已生成，等待人工审批", 15.2)

    add_heading(doc, "6.3 数据分析页面", 2)
    add_para(doc, "数据分析页面把历史需求与未来预测拆成两个图表：历史曲线使用 30 天滚动累计，未来预测单独用柱状图展示累计需求和缺口，避免单日需求与累计值共用坐标轴造成误读。")
    add_picture(doc, SCREENSHOTS / "03-analysis-charts.png", "图 6-3 数据分析页：历史 30 天滚动累计与未来 30 天预测/缺口", 15.2)

    add_heading(doc, "6.4 审批与草稿编辑", 2)
    add_para(doc, "采购审批页展示订单状态、金额、整单预算和正式单号。采购员可以查看明细，并对草稿修改数量和整单预算；服务端重新校验预测证据、推荐证据、产能、MOQ 和预算。")
    add_picture(doc, SCREENSHOTS / "04-approval-list.png", "图 6-4 采购审批列表", 15.2)
    add_picture(doc, SCREENSHOTS / "05-approval-detail.png", "图 6-5 采购订单明细：物料、单位、预算、交付要求和推荐依据", 15.2)
    add_picture(doc, SCREENSHOTS / "06-draft-edit.png", "图 6-6 草稿编辑：采购员可修改数量和整单预算", 15.2)

    add_heading(doc, "6.5 执行记录与工具追踪", 2)
    add_para(doc, "执行记录页保存每次 Agent run 的状态、业务结论、关联订单和工具调用 Trace。排查详情可以展开查看每个工具的参数和输出，便于复现实验和定位问题。")
    add_picture(doc, SCREENSHOTS / "07-execution-records.png", "图 6-7 执行记录：历史任务与业务结论", 15.2)
    add_picture(doc, SCREENSHOTS / "08-execution-traces.png", "图 6-8 执行记录：Agent 工具调用追踪", 15.2)
    page_break(doc)

    # 7
    add_heading(doc, "7 实验结果分析", 1)
    add_heading(doc, "7.1 基线预测结果", 2)
    demand_metrics = baseline["demand_metrics"]
    shortage_metrics = baseline["shortage_metrics"]
    add_table(doc, ["指标", "需求预测", "缺口预测"], [
        ["样本数", demand_metrics["sample_count"], shortage_metrics["sample_count"]],
        ["MAPE", f'{demand_metrics["mape_percent"]}%', f'{shortage_metrics["mape_percent"]}%'],
        ["MAE", demand_metrics["mae"], shortage_metrics["mae"]],
        ["WAPE", f'{demand_metrics["wape_percent"]}%', f'{shortage_metrics["wape_percent"]}%'],
        ["MAPE 覆盖", f'{demand_metrics["mape_coverage_percent"]}%', f'{shortage_metrics["mape_coverage_percent"]}%'],
    ], widths=[4.0, 5.5, 5.5])
    add_para(doc, "基线在 160 个 30 天窗口上的需求和缺口 MAPE 均低于 8%，说明 30 天移动平均对当前模拟数据是稳定基线。但它不能表达趋势和季节性，因此仍需与 Prophet 比较。")

    add_heading(doc, "7.2 预测模型对比", 2)
    rows = []
    for key, label in [
        ("moving_average_30d", "30 天移动平均"),
        ("prophet", "长历史 Prophet"),
        ("prophet_dynamic_seasonality", "动态 Prophet"),
        ("hybrid_dynamic_prophet_or_moving_average_30d", "混合策略"),
    ]:
        item = forecast[key]
        rows.append([
            label,
            item.get("status", "-"),
            item.get("demand_metrics", {}).get("mape_percent", "-"),
            item.get("shortage_metrics", {}).get("mape_percent", "-"),
            item.get("success_rate", {}).get("successful_windows", "-"),
            item.get("success_rate", {}).get("total_windows", "-"),
        ])
    add_table(doc, ["模型", "状态", "需求 MAPE", "缺口 MAPE", "成功窗口", "总窗口"], rows, widths=[4.2, 2.0, 2.4, 2.4, 2.2, 2.0])
    add_picture(doc, CHARTS / "forecast_mape.png", "图 7-1 预测模型滚动回测 MAPE 对比", 15.0)
    add_para(doc, "长历史 Prophet 只在 50/160 个窗口成功，主要原因是短训练历史；动态 Prophet 成功 140/160，缺口 MAPE 5.34%；混合策略在 20 个窗口回退移动平均，覆盖 160/160，缺口 MAPE 5.72%。因此正确表述是“混合策略覆盖全部窗口”，不能表述为“纯 Prophet 全部达标”。")

    add_heading(doc, "7.3 供应商推荐结果", 2)
    add_table(doc, ["数据集", "案例数", "综合评分 Top-3", "价格排序 Top-3", "标签说明"], [
        ["冻结集", supplier["frozen"]["case_count"], f'{supplier["frozen"]["composite_top3_hit_rate_percent"]}%', f'{supplier["frozen"]["price_sort_top3_hit_rate_percent"]}%', "场景阈值规则标签"],
        ["挑战集", supplier["challenge"]["case_count"], f'{supplier["challenge"]["composite_top3_hit_rate_percent"]}%', f'{supplier["challenge"]["price_sort_top3_hit_rate_percent"]}%', "场景阈值规则标签"],
    ], widths=[3.0, 2.0, 4.2, 4.2, 4.0])
    add_picture(doc, CHARTS / "supplier_hit.png", "图 7-2 供应商推荐 Top-3 命中率对比", 14.0)
    add_para(doc, "综合评分在两个数据集上均达到 100%，价格排序基线为 93.33%/95.00%。该结果说明“不是只看最低价”在当前模拟场景中更稳定。但命中定义是“Top-3 中至少出现一个可接受供应商”，标签由模拟场景阈值生成，不能等同于真实人工标签或严格的最佳供应商排名。")

    add_heading(doc, "7.4 Agent 输入评测", 2)
    add_table(doc, ["类别", "用例数", "通过数", "通过率"], [
        [name, stats["total"], stats["passed"], f'{stats["pass_rate_percent"]}%']
        for name, stats in [
            ("正常采购", agent["by_category"]["normal"]),
            ("信息缺失", agent["by_category"]["missing"]),
            ("条件校验", agent["by_category"]["validation"]),
            ("意图判断", agent["by_category"]["intent"]),
            ("语言/单位", agent["by_category"]["language"]),
            ("安全越权", agent["by_category"]["security"]),
        ]
    ], widths=[4.0, 3.5, 3.5, 3.5])
    add_picture(doc, CHARTS / "agent_eval.png", "图 7-3 Agent 输入评测分类通过情况", 15.0)
    add_para(doc, f'本次真实模型评测共 {agent["case_count"]} 条，通过 {agent["passed"]} 条，通过率 {agent["pass_rate_percent"]}%；平均耗时 {agent["avg_latency_ms"] / 1000:.1f} 秒，P95 耗时 {agent["p95_latency_ms"] / 1000:.1f} 秒。用例覆盖缺数量/预算/交期、物料歧义、预算冲突、预算不足、交期不足、产能不足、MOQ、只分析、条件式采购、单位不一致、越权审批和预算覆盖尝试。')

    add_heading(doc, "7.5 审批与 MySQL 集成", 2)
    add_para(doc, "MySQL 集成测试覆盖双会话并发审批、重复审批幂等、事务回滚和重连后的 Agent run/trace/order 可查询性。证据中并发审批只有一个成功、一个冲突，审批记录为 1 条；重复审批复用同一订单；回滚后订单仍为待审批且没有审批记录。")
    add_para(doc, "审批安全边界包括：buyer/approver 角色校验、订单版本校验、报价/产能/MOQ/金额/预算复核、正式单号生成和审计日志。模型没有审批工具，不能通过自然语言直接批准订单。")
    page_break(doc)

    # 8
    add_heading(doc, "8 结论与未来工作", 1)
    add_heading(doc, "8.1 结论", 2)
    add_para(doc, "本项目完成了从模拟 ERP 数据、库存与缺口计算、混合需求预测、供应商推荐、采购草稿、Agent 工具编排到人工审批的完整闭环。后端掌握 SQL、计算、评分、预算、权限、事务和幂等，LLM 只负责理解与编排，符合课程对“算法研究类”和“业务安全”的要求。")
    add_para(doc, "实验表明混合预测在固定库存滚动回测下缺口 MAPE 为 5.72%，动态 Prophet 在成功窗口为 5.34%；供应商综合评分 Top-3 达到 100%；26 条 Agent 输入用例全部通过；MySQL 并发、幂等和回滚测试通过。系统具备可演示、可复现和可解释的证据链。")
    add_heading(doc, "8.2 局限", 2)
    add_bullets(doc, [
        "预测使用固定 2026-10-10 库存快照，且测试窗口已被诊断复看，不能表述为完全独立的最终测试。",
        "长历史 Prophet 仅覆盖 50/160 窗口；混合策略包含 20 次移动平均回退，不是纯 Prophet。",
        "供应商标签由模拟场景阈值生成，Top-3 命中口径较宽松，不是真实人工标签。",
        "当前别名表是演示数据，换库仍需导入别名；更通用的方向是候选集 + 模型选 ID + 人工确认。",
        "对象级读权限、前端代码分割、历史库存回放和更大规模 Agent 评测仍有改进空间。",
    ])
    add_heading(doc, "8.3 未来工作", 2)
    add_bullets(doc, [
        "将物料解析改为“后端候选集 + 模型语义选 ID + 人工确认”，别名退化为可选缓存。",
        "增加历史库存快照和在途采购，做真正的历史缺口回放。",
        "增加独立最终留出集、供应商精确排名指标（MRR/NDCG）和更大规模 Agent 任务集。",
        "补充英文/多语言、单位换算、模糊日期和异常值处理。",
        "完善对象级权限、前端懒加载和生产部署配置。",
    ])
    page_break(doc)

    # Appendix
    add_heading(doc, "附录 A 复现命令与项目结构", 1)
    add_heading(doc, "A.1 项目结构", 2)
    add_code(doc, "backend/app/          FastAPI、SQLAlchemy 模型、服务、算法、Agent\nbackend/alembic/      数据库迁移\nbackend/tests/        单元/集成/安全测试\nbackend/evaluation/   预测、推荐、Agent 评测脚本\nfrontend/src/         Vue 3 四页面工作台\ndata/raw/             模拟数据 CSV 与元数据\nreports/generated/    评测 JSON、图表、截图（Git 忽略）\ndocs/                 需求、架构、API、数据字典、验收文档")
    add_heading(doc, "A.2 复现命令", 2)
    add_code(doc, "python -m venv .venv\n.\\.venv\\Scripts\\python -m pip install -r requirements.txt\npython -m alembic upgrade head\npython -m backend.app.db.init_demo_data --database-url \"sqlite+pysqlite:///./erp_demo.sqlite\"\npython -m uvicorn backend.app.main:app --host 127.0.0.1 --port 8000\nnpm install --prefix frontend\nnpm run dev --prefix frontend")
    add_heading(doc, "A.3 实验证据文件", 2)
    add_bullets(doc, [
        "reports/generated/baseline_eval.json",
        "reports/generated/supplier_eval.json",
        "reports/generated/forecast_comparison.json",
        "reports/generated/forecast_diagnostics/",
        "reports/generated/agent_input_eval.json",
        "reports/generated/mysql_integration.json",
        "reports/generated/screenshots_current/",
        "reports/generated/report_charts/",
    ])
    add_para(doc, "说明：本报告中的姓名、学号、班级为待填写占位符；提交前请按课程要求将文件夹和 Word 文件名改为实际组员姓名。")

    DELIVERY.mkdir(parents=True, exist_ok=True)
    output = DELIVERY / DOCX_NAME
    doc.save(output)
    return output


def ignore_patterns(_dir, names):
    ignored = set()
    for name in names:
        if name in {".git", ".venv", "node_modules", "dist", "__pycache__", ".pytest_cache", ".mypy_cache", ".ruff_cache", "delivery"}:
            ignored.add(name)
        elif name.endswith((".pyc", ".sqlite", ".db", ".tsbuildinfo")):
            ignored.add(name)
        elif name.startswith("chrome-profile"):
            ignored.add(name)
    return ignored


def copy_package() -> Path:
    package = DELIVERY / "源代码和数据集"
    if package.exists():
        shutil.rmtree(package)
    package.mkdir(parents=True)

    for name in ("backend", "frontend", "scripts", "docs", "data"):
        src = ROOT / name
        if src.exists():
            shutil.copytree(src, package / name, ignore=ignore_patterns)

    for name in ("requirements.txt", "alembic.ini", "docker-compose.yml", "README.md", "AGENTS.md", ".env.example"):
        src = ROOT / name
        if src.exists():
            shutil.copy2(src, package / name)

    evidence = package / "reports" / "generated"
    evidence.mkdir(parents=True, exist_ok=True)
    for name in (
        "baseline_eval.json", "supplier_eval.json", "forecast_comparison.json",
        "agent_input_eval.json", "agent_input_eval_cases.csv", "mysql_integration.json",
    ):
        src = EVIDENCE / name
        if src.exists():
            shutil.copy2(src, evidence / name)
    for directory in ("forecast_diagnostics", "screenshots_current", "report_charts"):
        src = EVIDENCE / directory
        if src.exists():
            shutil.copytree(src, evidence / directory, ignore=ignore_patterns)
    return package


if __name__ == "__main__":
    report = build_report()
    package = copy_package()
    print(f"report: {report}")
    print(f"package: {package}")
