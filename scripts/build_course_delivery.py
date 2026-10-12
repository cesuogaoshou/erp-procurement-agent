"""Build the course Word report and delivery folder from current evidence files."""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

try:
    import matplotlib.pyplot as plt
except Exception:  # pragma: no cover - surfaced in CLI failure
    plt = None


ROOT = Path(__file__).resolve().parents[1]
DELIVERY_ROOT = ROOT / "delivery"
PACKAGE_NAME = "【待填写姓名】ERP课程作业报告"
REPORT_NAME = f"{PACKAGE_NAME}.docx"
SOURCE_DIR_NAME = "源代码和数据集"
PLACEHOLDER = "【待填写：姓名 / 学号 / 班级】"


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8", errors="replace"), strict=False)


def pct(value: Any) -> str:
    return f"{value}%"


def git_text(*args: str) -> str:
    try:
        return subprocess.check_output(["git", *args], cwd=ROOT, text=True, stderr=subprocess.DEVNULL).strip()
    except Exception:
        return "unavailable"


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, fill: str | None = None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_font(run, size=9.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None, trailing_paragraph: bool = True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, fill="F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    if trailing_paragraph:
        doc.add_paragraph()
    return table


def add_para(doc: Document, text: str = "", bold: bool = False, size: float = 11, color: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_font(r, size=10.5)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=9.5, bold=True, color="555555")


def add_picture(doc: Document, path: Path, caption: str, width: float = 5.9):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)
    else:
        add_para(doc, f"[缺失图片：{path.name}]", bold=True, color="9B1C1C")


def style_doc(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ERP智能采购Agent课程报告")
    set_font(run, size=9, color="666666")


def extract_evidence() -> dict[str, Any]:
    baseline = read_json(ROOT / "reports/generated/baseline_eval.json")
    supplier = read_json(ROOT / "reports/generated/supplier_eval.json")
    forecast = read_json(ROOT / "reports/generated/forecast_comparison.json")
    diagnostics = read_json(ROOT / "reports/generated/forecast_diagnostics/forecast_diagnostics.json")
    real = read_json(ROOT / "reports/generated/real_model_acceptance.json")
    mysql = read_json(ROOT / "reports/generated/mysql_integration.json")
    metadata = read_json(ROOT / "data/raw/generation_metadata.json")
    models = forecast["models"]
    return {
        "baseline": baseline,
        "supplier": supplier,
        "forecast": forecast,
        "diagnostics": diagnostics,
        "real": real,
        "mysql": mysql,
        "metadata": metadata,
        "forecast_rows": [
            ("移动平均基线", "160/160", models["moving_average_30d"]["demand_metrics"]["mape_percent"], models["moving_average_30d"]["shortage_metrics"]["mape_percent"], "全样本固定库存情景"),
            ("长历史 Prophet", f"{models['prophet']['success_rate']['successful_windows']}/160", models["prophet"]["demand_metrics"]["mape_percent"], models["prophet"]["shortage_metrics"]["mape_percent"], "仅成功窗口，失败窗口单列"),
            ("同窗口移动平均", f"{models['moving_average_30d_on_prophet_success_windows']['success_rate']['successful_windows']}/160", models["moving_average_30d_on_prophet_success_windows"]["demand_metrics"]["mape_percent"], models["moving_average_30d_on_prophet_success_windows"]["shortage_metrics"]["mape_percent"], "与长历史 Prophet 成功窗口对齐"),
            ("动态 Prophet", f"{models['prophet_dynamic_seasonality']['success_rate']['successful_windows']}/160", models["prophet_dynamic_seasonality"]["demand_metrics"]["mape_percent"], models["prophet_dynamic_seasonality"]["shortage_metrics"]["mape_percent"], "短历史周季节，长历史年+周季节"),
            ("混合服务策略", "160/160", models["hybrid_dynamic_prophet_or_moving_average_30d"]["demand_metrics"]["mape_percent"], models["hybrid_dynamic_prophet_or_moving_average_30d"]["shortage_metrics"]["mape_percent"], "动态 Prophet 或 30日移动平均回退"),
        ],
        "supplier_rows": [
            ("冻结案例", supplier["frozen"]["case_count"], supplier["frozen"]["price_sort_top3_hit_rate_percent"], supplier["frozen"]["composite_top3_hit_rate_percent"], f"{supplier['frozen']['label_width']['min']} - {supplier['frozen']['label_width']['max']}"),
            ("挑战案例", supplier["challenge"]["case_count"], supplier["challenge"]["price_sort_top3_hit_rate_percent"], supplier["challenge"]["composite_top3_hit_rate_percent"], f"{supplier['challenge']['label_width']['min']} - {supplier['challenge']['label_width']['max']}"),
        ],
    }


def make_charts(out_dir: Path, ev: dict[str, Any]) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = {}
    if plt is not None:
        labels = [r[0] for r in ev["forecast_rows"]]
        demand = [float(r[2]) for r in ev["forecast_rows"]]
        shortage = [float(r[3]) for r in ev["forecast_rows"]]
        x = range(len(labels))
        plt.figure(figsize=(8, 4.5), dpi=180)
        plt.bar([i - 0.18 for i in x], demand, width=0.36, label="需求MAPE")
        plt.bar([i + 0.18 for i in x], shortage, width=0.36, label="缺口MAPE")
        plt.axhline(8, color="#9B1C1C", linestyle="--", linewidth=1, label="8%参考线")
        plt.xticks(list(x), labels, rotation=20, ha="right")
        plt.ylabel("MAPE (%)")
        plt.title("预测实验指标（探索性固定库存情景）")
        plt.legend()
        plt.tight_layout()
        p = out_dir / "forecast_metrics.png"
        plt.savefig(p)
        plt.close()
        charts["forecast"] = p

        labels = [r[0] for r in ev["supplier_rows"]]
        price = [float(r[2]) for r in ev["supplier_rows"]]
        comp = [float(r[3]) for r in ev["supplier_rows"]]
        x = range(len(labels))
        plt.figure(figsize=(6.8, 3.8), dpi=180)
        plt.bar([i - 0.18 for i in x], price, width=0.36, label="价格排序Top-3")
        plt.bar([i + 0.18 for i in x], comp, width=0.36, label="四指标综合Top-3")
        plt.ylim(80, 102)
        plt.xticks(list(x), labels)
        plt.ylabel("Top-3命中率 (%)")
        plt.title("供应商推荐评测")
        plt.legend()
        plt.tight_layout()
        p = out_dir / "supplier_metrics.png"
        plt.savefig(p)
        plt.close()
        charts["supplier"] = p
    else:
        def simple_grouped_bars(path: Path, title: str, labels: list[str], a: list[float], b: list[float], a_label: str, b_label: str, y_min: float = 0.0, y_max: float | None = None):
            y_max = y_max or max(max(a), max(b), 8.0) * 1.15
            img = Image.new("RGB", (1400, 780), "white")
            draw = ImageDraw.Draw(img)
            try:
                font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 34)
                small = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 24)
            except Exception:
                font = ImageFont.load_default()
                small = ImageFont.load_default()
            left, top, right, bottom = 110, 100, 1320, 610
            draw.text((left, 35), title, fill="#0B2545", font=font)
            draw.line((left, bottom, right, bottom), fill="#333333", width=2)
            draw.line((left, top, left, bottom), fill="#333333", width=2)
            group_w = (right - left) / len(labels)
            scale = (bottom - top) / (y_max - y_min)
            for i, label in enumerate(labels):
                cx = left + i * group_w + group_w / 2
                vals = [(a[i], "#2E74B5"), (b[i], "#7A5A00")]
                for j, (val, color) in enumerate(vals):
                    x1 = cx - 42 + j * 44
                    x2 = x1 + 34
                    y1 = bottom - (val - y_min) * scale
                    draw.rectangle((x1, y1, x2, bottom), fill=color)
                    draw.text((x1 - 8, y1 - 30), f"{val:.2f}", fill="#333333", font=small)
                draw.multiline_text((cx - group_w / 2 + 10, bottom + 18), label.replace(" ", "\n"), fill="#333333", font=small, spacing=2)
            draw.rectangle((left, 665, left + 28, 693), fill="#2E74B5")
            draw.text((left + 38, 662), a_label, fill="#333333", font=small)
            draw.rectangle((left + 310, 665, left + 338, 693), fill="#7A5A00")
            draw.text((left + 348, 662), b_label, fill="#333333", font=small)
            img.save(path)

        labels = [r[0] for r in ev["forecast_rows"]]
        demand = [float(r[2]) for r in ev["forecast_rows"]]
        shortage = [float(r[3]) for r in ev["forecast_rows"]]
        p = out_dir / "forecast_metrics.png"
        simple_grouped_bars(p, "预测实验指标（探索性固定库存情景）", labels, demand, shortage, "需求MAPE", "缺口MAPE", 0, 14)
        charts["forecast"] = p

        labels = [r[0] for r in ev["supplier_rows"]]
        price = [float(r[2]) for r in ev["supplier_rows"]]
        comp = [float(r[3]) for r in ev["supplier_rows"]]
        p = out_dir / "supplier_metrics.png"
        simple_grouped_bars(p, "供应商推荐评测", labels, price, comp, "价格排序Top-3", "四指标Top-3", 80, 102)
        charts["supplier"] = p

    arch = out_dir / "architecture_flow.png"
    img = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 34)
        small = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 24)
    except Exception:
        font = ImageFont.load_default()
        small = ImageFont.load_default()
    boxes = [
        (70, 120, 330, 240, "Vue演示前端\n采购/分析/审批/记录"),
        (430, 120, 700, 240, "FastAPI API层\n认证与路由"),
        (800, 80, 1130, 210, "Agent Loop\nFunction Calling编排"),
        (800, 260, 1130, 390, "业务服务层\n库存/预测/推荐/草稿/审批"),
        (1230, 120, 1510, 240, "MySQL/SQLite\nERP模拟数据与Trace"),
        (430, 520, 700, 660, "人工审批\n角色权限/幂等/审计"),
        (800, 520, 1130, 660, "评测与证据\nJSON/CSV/曲线/截图"),
    ]
    for x1, y1, x2, y2, text in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, outline="#2E74B5", width=4, fill="#F4F6F9")
        draw.multiline_text((x1 + 24, y1 + 26), text, fill="#0B2545", font=font, spacing=8)
    arrows = [
        ((330, 180), (430, 180)), ((700, 180), (800, 150)), ((700, 180), (800, 325)),
        ((1130, 150), (1230, 180)), ((1130, 325), (1230, 180)), ((565, 240), (565, 520)),
        ((700, 590), (800, 590)), ((965, 390), (965, 520)),
    ]
    for start, end in arrows:
        draw.line((*start, *end), fill="#1F4D78", width=4)
        ex, ey = end
        draw.polygon([(ex, ey), (ex - 14, ey - 8), (ex - 14, ey + 8)], fill="#1F4D78")
    draw.text((70, 40), "ERP智能采购Agent系统架构与数据流", fill="#0B2545", font=font)
    draw.text((70, 805), "边界：LLM仅负责意图理解与工具编排；SQL、计算、评分、审批、事务与幂等由后端服务负责。", fill="#555555", font=small)
    img.save(arch)
    charts["architecture"] = arch
    return charts


def build_report(package_dir: Path, ev: dict[str, Any], charts: dict[str, Path]) -> Path:
    doc = Document()
    style_doc(doc)
    report_path = package_dir / REPORT_NAME
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("基于 LLM Agent 的 ERP 智能采购决策系统")
    set_font(r, size=22, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(subtitle.add_run("课程选题 3.2：AI Agent 驱动的采购订单自动生成与供应商推荐算法研究"), size=12, color="555555")
    add_para(doc, PLACEHOLDER, bold=True, color="9B1C1C")
    add_para(doc, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}；代码基线：{git_text('rev-parse', '--short', 'HEAD')}；参考基线：bc45878。")
    add_para(doc, "说明：姓名、学号、班级无法从项目资料确认，以上保留醒目占位。报告中的实验数值均由当前文件系统中的原始 JSON/CSV/PNG 证据读取，不使用聊天摘要。", bold=True)
    doc.add_page_break()

    doc.add_heading("1. 项目背景与采购业务需求", level=1)
    add_para(doc, "本项目面向课程 3.2，构建一个 ERP 智能采购决策系统。核心业务是采购员用自然语言提出“分析未来30天物料缺口，推荐供应商并生成采购草稿”，系统完成库存查询、需求预测、缺口计算、供应商比较、草稿生成和人工审批。")
    add_para(doc, "项目边界保持在课程要求内：不扩展为完整 ERP，不接入真实供应商、支付或邮件，不建设 RAG/MCP/多 Agent 平台。LLM 的责任被限制为意图理解和工具编排，所有 SQL、金额、评分、业务校验、审批、事务和幂等都由后端服务完成。")
    add_bullets(doc, [
        "采购员：发起智能采购、提交草稿、查看自己的运行和订单。",
        "审批员：查看待审批订单，基于后端证据批准或驳回。",
        "系统：保存 Agent run、工具 trace、采购订单、审批记录和审计日志。",
    ])

    doc.add_heading("2. 模拟数据定义与预处理", level=1)
    meta = ev["metadata"]
    add_para(doc, f"数据集为 `{meta['dataset']}`，固定随机种子 {meta['seed']}，基准日期 {meta['base_date']}，包含 {meta['materials']} 种物料、{meta['suppliers']} 家供应商和 {meta['history_days']} 天逐日需求历史。需求模式为趋势、周季节、年度波动、噪声和间歇性尖峰的组合。")
    add_table(doc, ["CSV/证据", "用途", "预处理说明"], [
        ["materials.csv / suppliers.csv", "物料与供应商主数据", "编码、名称、启用状态等基础字段"],
        ["inventory_snapshots.csv", "库存快照", f"静态快照日期 {meta['inventory_snapshot_date']}；报告中按固定库存情景解释"],
        ["demand_history.csv", "需求预测与回测", "按物料和日期聚合，滚动窗口训练与30天持出"],
        ["supplier_materials.csv", "报价、交期、MOQ、容量", "推荐前进行关系、容量、预算、交期过滤"],
        ["supplier_performance.csv", "价格、交付、质检、投诉", "归一化后进入四指标综合评分"],
    ], widths=[1.9, 1.7, 2.7])
    add_para(doc, "预处理过程避免把未来信息交给决策输入。需要特别说明的是，当前数据没有历史库存流水，因此缺口回测不是历史库存快照回放，而是在同一个 2026-10-10 静态库存快照下进行的固定库存情景比较。")

    doc.add_heading("3. 系统架构、数据库及实现步骤", level=1)
    add_picture(doc, charts["architecture"], "图1 系统架构与数据流：LLM 只编排工具，后端拥有业务边界。", 6.2)
    add_para(doc, "后端采用 FastAPI、Pydantic、SQLAlchemy、Alembic 和 pytest；目标集成数据库为 MySQL，单元测试可使用 SQLite。前端采用 Vue 3、TypeScript、Element Plus 和 ECharts，提供智能采购、数据分析、采购审批、执行记录四个页面。")
    add_table(doc, ["步骤", "实现内容", "关键证据"], [
        ["1", "建立数据模型、迁移和模拟种子", "backend/alembic、data/raw、docs/DATA_DICTIONARY.md"],
        ["2", "实现库存、预测、推荐、草稿与审批服务", "backend/app/services、backend/app/algorithms"],
        ["3", "实现真实 Function Calling Agent Loop", "agent_loop.py、agent_provider.py、agent_tools.py"],
        ["4", "接入人工审批、权限、事务与审计", "approvals.py、auth.py、MySQL 集成报告"],
        ["5", "生成评测与前端演示证据", "reports/generated、frontend/src/App.vue"],
    ], widths=[0.6, 2.4, 3.2])

    doc.add_heading("4. Prophet 与混合预测策略", level=1)
    add_para(doc, "预测目标是未来30天累计需求及由此计算的物料缺口。当前服务策略冻结为 `forecast_strategy_v1_20261011`：以动态 Prophet 为主，短历史或 Prophet 失败时回退到 30 日移动平均。报告中的预测实验属于探索性、固定库存情景；不宣称独立、未触碰的最终测试集达标。")
    add_table(doc, ["模型/策略", "窗口覆盖", "需求MAPE", "缺口MAPE", "说明"], ev["forecast_rows"], widths=[1.3, 1.0, 0.9, 0.9, 2.2])
    add_picture(doc, charts["forecast"], "图2 预测实验指标对比：纯 Prophet 成功窗口与混合全样本结果分开报告。", 5.8)
    diag = ev["diagnostics"]
    add_para(doc, f"诊断检查覆盖 {diag['implementation_checks']['window_count']} 个窗口和 {diag['implementation_checks']['demand_rows_checked']} 行需求数据，未发现日期排序、持出长度、日连续性、负预测截断或缺口公式复用问题。诊断同时记录：测试区间已被检查，且 30 天步长只剩 {diag['split_usage']['unused_tail_days_for_30d_stride']['start']} 至 {diag['split_usage']['unused_tail_days_for_30d_stride']['end']} 的尾部不足以形成新的30天持出窗口。")
    add_picture(doc, ROOT / "reports/generated/forecast_diagnostics/curves/material_1_cutoff_2026-06-27.png", "图3 Prophet 代表性预测曲线（当前版本诊断输出）。", 5.8)

    doc.add_heading("5. 四指标供应商推荐及自主贡献", level=1)
    add_para(doc, "供应商推荐先进行硬约束过滤，再按四指标综合评分排序。四指标为历史价格、准时交付率、质量合格率和投诉表现，默认权重为价格0.35、交付0.30、质量0.25、投诉0.10。价格和投诉是成本型指标，交付和质量是收益型指标。")
    add_table(doc, ["案例集", "案例数", "价格排序Top-3", "四指标Top-3", "标签宽度"], ev["supplier_rows"], widths=[1.3, 0.8, 1.2, 1.2, 1.0])
    add_picture(doc, charts["supplier"], "图4 供应商推荐 Top-3 命中率：Top-3 指至少一个可接受供应商进入前三。", 5.5)
    add_para(doc, f"Top-3 的定义来自当前 `supplier_eval.json`：{ev['supplier']['frozen']['top3_definition']}。标签来源为模拟场景阈值规则，冻结案例标签宽度 {ev['supplier']['frozen']['label_width']}，挑战案例标签宽度 {ev['supplier']['challenge']['label_width']}。这些标签不是被评测的同一加权公式，也不是企业真实人工标签，因此报告只说明其模拟来源和宽度。")
    add_para(doc, "自主贡献主要体现在自建 ERP 采购模拟数据、场景阈值标签、约束过滤加四指标综合评分、MOQ/包装倍数采购草稿校验，以及将真实 Function Calling 与后端审批安全边界结合。")

    doc.add_heading("6. Function Calling 与审批闭环", level=1)
    real = ev["real"]
    add_para(doc, f"当前真实模型验收报告状态为 `{real['status']}`，模型为 `{real['model_used']}`，`base_url_configured=true`，并记录了非密钥配置来源。happy path 首次运行产生采购草稿 order_id={real['happy_path']['first']['purchase_order_id']}，工具序列为 {', '.join(real['happy_path']['first']['tool_calls'])}。")
    add_table(doc, ["环节", "证据字段", "当前结果"], [
        ["草稿创建", "happy_path.first.status", real["happy_path"]["first"]["status"]],
        ["重复请求", "happy_path.duplicate_reused_order", real["happy_path"]["duplicate_reused_order"]],
        ["模型审批尝试", "approval_loop.model_approval_attempt.tool_calls", real["approval_loop"]["model_approval_attempt"]["tool_calls"]],
        ["人工审批", "approval_loop.human_approval_status", real["approval_loop"]["human_approval_status"]],
        ["正式订单", "approval_loop.formal_order_number", real["approval_loop"]["formal_order_number"]],
        ["边界案例", "boundary_cases", ", ".join(f"{k}:{v['status']}" for k, v in real["boundary_cases"].items())],
    ], widths=[1.2, 2.5, 2.4])
    mysql = ev["mysql"]
    add_para(doc, f"MySQL 集成证据显示，并发审批中一个请求通过、另一个返回 `{mysql['core']['approval_concurrency']['outcomes'][1]['error']}`；重复审批只保留 {mysql['duplicate_approval']['approval_records']} 条审批记录；回滚场景保持订单为 `{mysql['core']['rollback']['status']}`。")

    doc.add_heading("7. 实验截图与结果分析", level=1)
    add_picture(doc, ROOT / "reports/generated/screenshots/01-buyer-agent-home.png", "图5 采购员智能采购页：自然语言触发真实后端 Agent 流程。", 5.9)
    add_picture(doc, ROOT / "reports/generated/screenshots/02-analysis-page.png", "图6 数据分析页：展示预测、库存和模型策略信息。", 5.9)
    add_picture(doc, ROOT / "reports/generated/screenshots/03-approval-approved-order.png", "图7 审批页：人工审批后显示正式订单。", 5.9)
    add_para(doc, "结果表明，后端服务可以在课程演示范围内完成从自然语言采购意图到草稿、审批、正式订单的闭环；预测达到课程数值参考需要按混合策略和固定库存情景解释；供应商推荐在模拟标签口径下达到 Top-3 指标，但该指标是“任一可接受供应商进入前三”，不是单一最佳供应商精确命中。")

    doc.add_heading("8. 局限与改进方向", level=1)
    add_bullets(doc, [
        "预测实验不是未触碰的独立最终测试；诊断过程已经检查过部分持出区间。",
        "缺口回测使用静态库存快照，无法替代真实历史库存流水回放。",
        "供应商标签来自模拟阈值规则，不代表真实企业人工或履约标签。",
        "当前权限为课程演示级最小对象读访问，不是生产级多租户 ERP 权限体系。",
        "LangGraph 已用于图执行，但跨进程恢复依赖 SQL 持久化，不是独立 checkpoint 后端。",
        "后续可补历史库存流水、冻结新 holdout、增加 MRR/NDCG 等更严格推荐指标，并对前端做代码分割。",
    ])

    doc.add_heading("9. 复现命令与交付说明", level=1)
    add_table(doc, ["目的", "命令"], [
        ["后端测试", "python -m pytest backend/tests -q"],
        ["基线评测", "python -m backend.evaluation.run_baseline"],
        ["供应商评测", "python -m backend.evaluation.run_supplier_eval"],
        ["预测比较", "python -m backend.evaluation.run_forecast_comparison"],
        ["预测诊断", "python -m backend.evaluation.run_forecast_diagnostics"],
        ["真实模型验收", "python -m backend.evaluation.run_real_model_acceptance"],
        ["前端审计/构建", "npm --prefix frontend audit --json; npm --prefix frontend run build"],
    ], widths=[1.4, 4.8], trailing_paragraph=False)

    doc.save(report_path)
    return report_path


EXCLUDE_DIRS = {
    ".git", ".pytest_cache", "__pycache__", ".mypy_cache", ".ruff_cache", ".venv",
    "node_modules", "dist", "build", "delivery", "chrome-profile", "chrome-profile-2",
    "chrome-profile-3", "chrome-profile-inspect",
}
EXCLUDE_SUFFIXES = {".pyc", ".sqlite", ".db", ".pem", ".key", ".p12", ".crt"}
EXCLUDE_FILES = {".env"}


def should_copy(path: Path) -> bool:
    rel = path.relative_to(ROOT)
    parts = set(rel.parts)
    if parts & EXCLUDE_DIRS:
        return False
    if path.name in EXCLUDE_FILES:
        return False
    if path.name.startswith(".env") and path.name != ".env.example":
        return False
    if path.suffix.lower() in EXCLUDE_SUFFIXES:
        return False
    if "reports" in rel.parts and "generated" in rel.parts:
        allowed = path.suffix.lower() in {".json", ".csv", ".png"}
        return allowed and not (parts & {"chrome-profile", "chrome-profile-2", "chrome-profile-3", "chrome-profile-inspect"})
    return True


def copy_delivery(package_dir: Path):
    source_root = package_dir / SOURCE_DIR_NAME
    if source_root.exists():
        shutil.rmtree(source_root)
    source_root.mkdir(parents=True)
    copied = []
    for path in ROOT.rglob("*"):
        if not path.is_file() or not should_copy(path):
            continue
        rel = path.relative_to(ROOT)
        dst = source_root / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, dst)
        copied.append(rel.as_posix())
    return copied


SECRET_PATTERNS = [
    re.compile(r"(?i)(api[_-]?key|secret|token|password)\s*[:=]\s*['\"]?([A-Za-z0-9_\-]{16,})"),
    re.compile(r"sk-[A-Za-z0-9]{20,}"),
]


BENIGN_SECRET_MARKERS = (
    "replace-with-local-demo-secret",
    "<set locally, never commit>",
    "should-not-be-printed",
    "secret-value",
    "env-file-secret",
    "dev-demo-secret-change-me",
)


def scan_sensitive(package_dir: Path) -> tuple[list[str], list[str], int]:
    findings = []
    benign = []
    scanned = 0
    for path in (package_dir / SOURCE_DIR_NAME).rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".py", ".md", ".txt", ".json", ".csv", ".env", ".example", ".ini", ".toml", ".yaml", ".yml", ".ts", ".vue", ".js"}:
            continue
        scanned += 1
        text = path.read_text(encoding="utf-8", errors="ignore")
        for pat in SECRET_PATTERNS:
            if pat.search(text):
                rel = str(path.relative_to(package_dir))
                if any(marker in text for marker in BENIGN_SECRET_MARKERS):
                    benign.append(rel)
                else:
                    findings.append(rel)
                break
    return findings, benign, scanned


def write_package_docs(package_dir: Path, copied: list[str], findings: list[str], benign: list[str], scanned: int, report_path: Path):
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    manifest = package_dir / "PACKAGE_MANIFEST.md"
    manifest.write_text(
        "\n".join([
            "# 交付包清单",
            "",
            f"- 生成时间：{generated_at}",
            f"- 代码版本：{git_text('rev-parse', 'HEAD')}",
            f"- Git分支：{git_text('branch', '--show-current')}",
            f"- Git状态：{git_text('status', '--short') or 'clean'}",
            f"- Word报告：{report_path.name}",
            f"- 源代码和数据集目录：{SOURCE_DIR_NAME}",
            f"- 文件数量：{len(copied)}",
            "",
            "## 已纳入",
            "",
            "- 后端/前端源码、Alembic迁移、测试、评测脚本。",
            "- data/raw CSV和generation_metadata.json。",
            "- requirements.txt、frontend/package-lock.json、.env.example、README和docs说明。",
            "- reports/generated下的JSON/CSV/PNG评测证据、曲线和截图。",
            "",
            "## 已排除",
            "",
            "- .env、密钥、证书、虚拟环境、node_modules、__pycache__、pytest缓存、Chrome profile缓存、SQLite/DB快照、delivery目录。",
            "",
            "## 复现命令",
            "",
            "```powershell",
            "python -m pytest backend/tests -q",
            "python -m backend.evaluation.run_baseline",
            "python -m backend.evaluation.run_supplier_eval",
            "python -m backend.evaluation.run_forecast_comparison",
            "python -m backend.evaluation.run_forecast_diagnostics",
            "python -m backend.evaluation.run_real_model_acceptance",
            "npm --prefix frontend audit --json",
            "npm --prefix frontend run build",
            "$env:RUN_MYSQL_INTEGRATION='1'; $env:MYSQL_INTEGRATION_REPORT_PATH='reports/generated/mysql_integration.json'; python -m pytest backend/tests/test_mysql_integration.py -q",
            "```",
        ]) + "\n",
        encoding="utf-8",
    )
    scan = package_dir / "SENSITIVE_SCAN.md"
    scan.write_text(
        "\n".join([
            "# 敏感信息检查",
            "",
            f"- 扫描时间：{generated_at}",
            f"- 扫描文本文件数：{scanned}",
            f"- 真实疑似密钥命中数量：{len(findings)}",
            f"- 占位/测试/演示常量命中数量：{len(benign)}",
            "",
            "## 真实疑似密钥",
            "",
            "未发现疑似明文密钥。" if not findings else "\n".join(f"- {x}" for x in findings),
            "",
            "## 已解释命中",
            "",
            "无。" if not benign else "\n".join(f"- {x}" for x in benign),
            "",
            "说明：脚本同时按复制规则排除了.env、密钥/证书文件、虚拟环境、node_modules、缓存和数据库快照。",
        ]) + "\n",
        encoding="utf-8",
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(DELIVERY_ROOT / PACKAGE_NAME))
    args = parser.parse_args()
    package_dir = Path(args.out).resolve()
    if package_dir.exists():
        shutil.rmtree(package_dir)
    package_dir.mkdir(parents=True)
    ev = extract_evidence()
    charts = make_charts(package_dir / "_generated_assets", ev)
    report_path = build_report(package_dir, ev, charts)
    copied = copy_delivery(package_dir)
    findings, benign, scanned = scan_sensitive(package_dir)
    write_package_docs(package_dir, copied, findings, benign, scanned, report_path)
    print(json.dumps({
        "package_dir": str(package_dir),
        "report_path": str(report_path),
        "copied_files": len(copied),
        "sensitive_findings": findings,
        "benign_sensitive_matches": benign,
        "generated_assets": {k: str(v) for k, v in charts.items()},
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
